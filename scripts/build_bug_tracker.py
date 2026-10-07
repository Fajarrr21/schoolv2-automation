#!/usr/bin/env python3
"""
Build Bug_Tracker_QA_CARDS.xlsx dari bugs.csv.

Pakai:
    python scripts/build_bug_tracker.py
    python scripts/build_bug_tracker.py bugs.csv Bug_Tracker_QA_CARDS.xlsx

bugs.csv adalah single source of truth (di-commit ke git).
File .xlsx adalah artifact hasil generate — regenerate kapan aja mau setor laporan.
"""
import csv
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC = sys.argv[1] if len(sys.argv) > 1 else "bugs.csv"
OUT = sys.argv[2] if len(sys.argv) > 2 else "Bug_Tracker_QA_CARDS.xlsx"

# --- palet warna (match template lama) ---
NAVY = "1F4E78"        # header bar + judul
SUBTLE = "595959"      # subtitle abu
PRIORITY_FILL = "D9E1F2"  # biru muda
SEV_FILL = {
    "High": "FFC7CE",     # merah muda
    "Medium": "FFEB9C",   # kuning
    "Low": "C6EFCE",      # hijau muda
}
STATUS_COLOR = {
    "Open": "C00000",     # merah
    "In Progress": "BF8F00",  # oranye gelap
    "Partially Fixed": "BF8F00",  # oranye gelap: sebagian sudah diperbaiki
    "Fixed": "548235",    # hijau
    "Closed": "548235",
    "Won't Fix": "808080",
}
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# kolom: (header, lebar)
COLS = [
    ("Bug ID", 10), ("Tiket", 20), ("Module", 18), ("Title", 34), ("Steps to Reproduce", 42),
    ("Expected Result", 42), ("Actual Result", 52), ("Severity", 11),
    ("Priority", 10), ("Reporter", 16), ("Status", 12), ("Date Found", 13),
    ("Video Evidence", 32), ("Screenshot Evidence", 32),
]
NAMES = [n for n, _ in COLS]
COL = {n: i for i, n in enumerate(NAMES, start=1)}  # nama kolom -> nomor kolom (1-based)
CENTER_COLS = {COL[n] for n in ("Bug ID", "Severity", "Priority", "Reporter", "Status", "Date Found")}
EVIDENCE_COLS = {COL["Video Evidence"], COL["Screenshot Evidence"]}  # hyperlink + warna
HYPERLINK_COLOR = "0563C1"           # biru standar Excel hyperlink

# Satu sheet per tiket (urutan tetap); tiket lain di csv ikut dibuatkan sheet di belakang.
TIKET_SHEET = {
    "Tiket 1 - Capaian Multi Select": "Tiket 1 - Multi Select",
    "Tiket 2 - Improvement Batas Tahfidz": "Tiket 2 - Batas Tahfidz",
    "Di luar tiket": "Di luar tiket",
}

with open(SRC, encoding="utf-8") as f:
    reader = csv.reader(f)
    header = next(reader)
    data = [row for row in reader if row and row[0].strip()]

if header[:len(NAMES)] != NAMES:
    sys.exit(f"Header {SRC} tidak sesuai urutan kolom: {NAMES}")


def tulis_sheet(ws, rows, judul):
    # baris 1: judul
    ws["A1"] = f"\U0001F41B BUG TRACKER \u2014 {judul}"
    ws["A1"].font = Font(name="Arial", bold=True, size=14, color=NAVY)

    # baris 2: subtitle + ringkasan
    sev = {s: sum(1 for r in rows if r[COL["Severity"] - 1] == s) for s in ("High", "Medium", "Low")}
    buka = sum(1 for r in rows if r[COL["Status"] - 1] not in ("Fixed", "Closed", "Won't Fix"))
    ws["A2"] = ("Reporter: Fajar Ardiansyah  |  Env: staging-school.cazh.id (V2)  |  "
                "Tester: QA Automation (Cypress)  |  "
                f"{len(rows)} bug ({buka} belum fixed): High {sev['High']}, "
                f"Medium {sev['Medium']}, Low {sev['Low']}")
    ws["A2"].font = Font(name="Arial", size=9, color=SUBTLE)

    # baris 4: header tabel
    hfill = PatternFill("solid", fgColor=NAVY)
    for ci, (name, width) in enumerate(COLS, start=1):
        c = ws.cell(row=HEADER_ROW, column=ci, value=name)
        c.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        c.fill = hfill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[chr(64 + ci)].width = width

    # data
    for ri, row in enumerate(rows, start=HEADER_ROW + 1):
        row = (row + [""] * len(COLS))[:len(COLS)]
        for ci, val in enumerate(row, start=1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(
                horizontal="center" if ci in CENTER_COLS else "left",
                vertical="top", wrap_text=True)
            c.border = BORDER
            # Evidence kolom (Video & Screenshot): jadikan hyperlink kalau ada path/value.
            # Format clickable: relative path -> di-buka pakai default app di OS.
            if ci in EVIDENCE_COLS and val and val.strip():
                c.hyperlink = val.strip()
                c.font = Font(name="Arial", size=10, color=HYPERLINK_COLOR, underline="single")
        # Severity
        sv = ws.cell(row=ri, column=COL["Severity"])
        fill = SEV_FILL.get((sv.value or "").strip())
        if fill:
            sv.fill = PatternFill("solid", fgColor=fill)
            sv.font = Font(name="Arial", size=10, bold=True)
        # Priority
        pr = ws.cell(row=ri, column=COL["Priority"])
        pr.fill = PatternFill("solid", fgColor=PRIORITY_FILL)
        pr.font = Font(name="Arial", size=10, bold=True)
        # Status
        st = ws.cell(row=ri, column=COL["Status"])
        sc = STATUS_COLOR.get((st.value or "").strip())
        if sc:
            st.font = Font(name="Arial", size=10, bold=True, color=sc)
        ws.row_dimensions[ri].height = 120

    ws.freeze_panes = f"A{HEADER_ROW + 1}"
    ws.sheet_view.showGridLines = False
    if rows:
        ws.auto_filter.ref = f"A{HEADER_ROW}:{chr(64 + len(COLS))}{HEADER_ROW + len(rows)}"


HEADER_ROW = 4
wb = Workbook()
wb.remove(wb.active)

tiket_list = list(TIKET_SHEET) + sorted({r[1] for r in data} - set(TIKET_SHEET))
for tiket in tiket_list:
    rows = [r for r in data if r[1] == tiket]
    if rows:
        tulis_sheet(wb.create_sheet(TIKET_SHEET.get(tiket, tiket)[:31]), rows, tiket)
tulis_sheet(wb.create_sheet("Semua Bug"), data, "Semua Tiket")

wb.save(OUT)
print(f"OK -> {OUT} ({len(data)} bugs)")
for ws in wb.worksheets:
    print(f"  {ws.title}: {ws.max_row - HEADER_ROW} bug")


# --- docs/evidence/README.md: indeks evidence per tiket (ikut di-generate) ---
import json
import os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TC_MAP = {}
for fx in ("tahfidz.json", "batas.json"):
    path = os.path.join(ROOT, "cypress", "fixtures", fx)
    if os.path.exists(path):
        for bug_id, info in (json.load(open(path, encoding="utf-8")).get("bug") or {}).items():
            if bug_id.startswith("BUG-"):
                TC_MAP[bug_id] = info.get("tc") or []

SEV_ORDER = {"High": 0, "Medium": 1, "Low": 2}
idx = [
    "# Evidence Bug — QA CARDS School V2 (staging)",
    "",
    "> Di-generate oleh `python scripts/build_bug_tracker.py` dari `bugs.csv` — jangan diedit manual.",
    "> Tiap folder `BUG-xxx/` berisi `README.md` (langkah, expected vs actual, lokasi source, saran fix) + screenshot.",
    "",
]
for tiket in tiket_list:
    rows = sorted((r for r in data if r[1] == tiket), key=lambda r: r[COL["Priority"] - 1])
    if not rows:
        continue
    terbuka = sum(1 for r in rows if r[COL["Status"] - 1] not in ("Fixed", "Closed", "Won't Fix"))
    sev = {s: sum(1 for r in rows if r[COL["Severity"] - 1] == s) for s in SEV_ORDER}
    idx += [
        f"## {tiket}",
        "",
        f"{len(rows)} bug · {terbuka} belum selesai · High {sev['High']} / Medium {sev['Medium']} / Low {sev['Low']}",
        "",
        "| Prioritas | Bug | Severity | Status | Judul | TC terkait | Evidence |",
        "|---|---|---|---|---|---|---|",
    ]
    for r in rows:
        bug = r[0]
        tc = ", ".join(TC_MAP.get(bug, [])) or "-"
        judul = r[COL["Title"] - 1].replace("|", "/")
        idx.append(
            f"| {r[COL['Priority'] - 1]} | {bug} | {r[COL['Severity'] - 1]} | {r[COL['Status'] - 1]} "
            f"| {judul} | {tc} | [{bug}/]({bug}/README.md) |"
        )
    idx.append("")
evid = os.path.join(ROOT, "docs", "evidence", "README.md")
open(evid, "w", encoding="utf-8").write("\n".join(idx))
print(f"OK -> docs/evidence/README.md (indeks per tiket)")
